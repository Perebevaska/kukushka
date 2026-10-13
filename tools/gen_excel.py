#!/usr/bin/env python3
"""Книга Excel финансового отдела «Зерна» для урока ch03/03 (синтетика, детерминированная).

    python3 tools/gen_excel.py          # нужен openpyxl: uv run --with openpyxl tools/gen_excel.py

Пишет в site/data/ (и блокнот site/notebooks/extra/ch03-excel.ipynb):
  zerno_finance.xlsx  — книга из трёх листов:
      «Описание» — пара строк текста о книге;
      «План»     — план выручки 2026 по кофейням и месяцам, тыс. ₽: две строки шапки
                   и пустая строка над таблицей, внизу строка «Итого», справа столбец «Год»;
      «Факт»     — факт за январь–сентябрь 2026 в том же виде, но шапка на строку длиннее
                   (дата выгрузки), справа «Итого» и текстовый «Комментарий»;
  finance_plan.csv,
  finance_fact.csv    — те же листы, сохранённые из Excel как «CSV (разделитель — точка с запятой)»:
                        шапка над таблицей сохраняется, дробная часть — через запятую.
                        в браузере .xlsx читается (openpyxl добавлен в локальный Pyodide); CSV-копии —
                        для задачи «тот же лист в CSV» и сверки plan-fact.

Заложено для урока:
  * над таблицей строки-заголовки: без skiprows/header первая строка станет «шапкой»;
  * у «Плана» и «Факта» разная высота шапки (3 и 4 строки);
  * строка «Итого» внизу: если её не убрать, сумма по столбцу удвоится;
  * EKB-1 открывается 15.10.2026: в плане январь–сентябрь пусто, факта нет;
  * KZN-2 в марте была закрыта на ремонт (провал факта, комментарий);
  * лишние для расчёта столбцы: «Кофейня», «Город», «Год»/«Итого», «Комментарий».
Файлы других уроков этот скрипт не трогает.
"""
import datetime as dt
from pathlib import Path

import numpy as np

OUT = Path(__file__).resolve().parent.parent / 'site' / 'data'
SEED = 2026_03
MONTHS = ['Янв', 'Фев', 'Мар', 'Апр', 'Май', 'Июн', 'Июл', 'Авг', 'Сен', 'Окт', 'Ноя', 'Дек']
FACT_MONTHS = MONTHS[:9]

# код, название, город, база плана в месяц (тыс. ₽), отношение факт/план, первый месяц работы (1–12)
SHOPS = [
    ('PRM-1', 'Ленина', 'Пермь', 1180, 1.03, 1),
    ('PRM-2', 'Компрос', 'Пермь', 760, 0.97, 1),
    ('KZN-1', 'Баумана', 'Казань', 1350, 1.05, 1),
    ('KZN-2', 'Кремлёвская', 'Казань', 820, 0.90, 1),
    ('TMN-1', 'Республики', 'Тюмень', 880, 0.99, 1),
    ('UFA-1', 'Ленина', 'Уфа', 700, 0.94, 1),
    ('EKB-1', 'Вайнера', 'Екатеринбург', 950, 1.0, 10),
]
# сезонность: зимой и осенью кофе пьют больше, летом меньше
SEASON = [1.00, 0.97, 1.02, 0.98, 0.95, 0.90, 0.86, 0.89, 1.03, 1.08, 1.10, 1.18]
COMMENTS = {'KZN-2': 'март: ремонт, закрыта 12 дней', 'UFA-1': 'вторая кофейня в городе не открылась — план завышен'}


def r1(x):
    return float(round(x + 0.0, 1))


def build():
    rng = np.random.default_rng(SEED)
    plan_rows, fact_rows = [], []
    for code, name, city, base, ratio, first in SHOPS:
        plan = []
        for m in range(1, 13):
            if m < first:
                plan.append(None)
            elif m == first and first == 10:          # открытие 15-го: полмесяца
                plan.append(r1(base * SEASON[m - 1] * 0.5))
            else:
                plan.append(r1(base * SEASON[m - 1] * rng.uniform(0.98, 1.02)))
        fact = []
        for m in range(1, 10):
            p = plan[m - 1]
            if p is None:
                fact.append(None)
                continue
            k = ratio * rng.uniform(0.96, 1.04)
            if code == 'KZN-2' and m == 3:
                k *= 0.6                             # ремонт
            fact.append(r1(p * k))
        plan_rows.append([code, name, city, *plan, r1(sum(v for v in plan if v is not None))])
        fact_rows.append([code, name, city, *fact, r1(sum(v for v in fact if v is not None)),
                          COMMENTS.get(code)])

    def total(rows, ncols):
        out = ['Итого', None, None]
        for j in range(3, 3 + ncols):
            vals = [r[j] for r in rows if r[j] is not None]
            out.append(r1(sum(vals)))
        return out

    plan_rows.append(total(plan_rows, 13))
    fact_rows.append(total(fact_rows, 10) + [None])
    plan_head = [['Сеть кофеен «Зерно»'], ['План выручки на 2026 год, тыс. ₽'], []]
    fact_head = [['Сеть кофеен «Зерно»'], ['Фактическая выручка 2026, тыс. ₽'],
                 ['Данные на 05.10.2026; сентябрь — предварительно'], []]
    plan = plan_head + [['Код', 'Кофейня', 'Город', *MONTHS, 'Год']] + plan_rows
    fact = fact_head + [['Код', 'Кофейня', 'Город', *FACT_MONTHS, 'Итого', 'Комментарий']] + fact_rows
    return plan, fact


def cell_csv(v):
    if v is None:
        return ''
    if isinstance(v, float):
        s = f'{v:.1f}'.replace('.', ',')
        return s
    v = str(v)
    return f'"{v}"' if (';' in v or '"' in v) else v


def write_csv(rows, path):
    width = max(len(r) for r in rows)
    lines = [';'.join(cell_csv(v) for v in list(r) + [None] * (width - len(r))) for r in rows]
    path.write_text('\n'.join(lines) + '\n', encoding='utf-8')


def write_xlsx(plan, fact, path):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = 'Описание'
    for line in ['Выручка сети кофеен «Зерно» за 2026 год',
                 'Листы: «План» — план по месяцам, «Факт» — факт январь–сентябрь.',
                 'Суммы — в тысячах рублей. Подготовил финансовый отдел.']:
        ws.append([line])
    ws['A1'].font = Font(bold=True)
    for title, rows in (('План', plan), ('Факт', fact)):
        ws = wb.create_sheet(title)
        for r in rows:
            ws.append(list(r))
        ws['A1'].font = Font(bold=True, size=13)
    stamp = dt.datetime(2026, 10, 5, 9, 0)
    wb.properties.created = stamp
    wb.properties.modified = stamp
    wb.properties.creator = 'Финансовый отдел «Зерна»'
    wb.save(path)


NB_OUT = OUT.parent / 'notebooks' / 'extra' / 'ch03-excel.ipynb'

NB_CELLS = [
    ('md', '''# Книга Excel: `read_excel` на своём компьютере

Блокнот к уроку «Типы: из текста в числа и даты» (глава 3), раздел «Книга Excel». В браузере курса те же примеры работают прямо в уроке; здесь — те же шаги в своём Jupyter на книге `zerno_finance.xlsx`, плюс запись результата обратно в Excel.

**Что нужно.** Библиотека openpyxl: в папке проекта из урока про uv выполните в терминале `uv add openpyxl` и перезапустите ядро (Kernel → Restart Kernel). Блокнот ищет книгу в папке `data` курса (`../../data`), а если не нашёл — рядом с собой.

**Как работать.** Выполняйте ячейки сверху вниз. В конце — три упражнения: место для кода помечено ✎, под каждым — ячейка самопроверки.'''),
    ('code', '''from pathlib import Path
import pandas as pd

DATA = Path("../../data")                   # блокнот лежит в notebooks/extra, данные — в data
if not (DATA / "zerno_finance.xlsx").exists():
    DATA = Path(".")                        # книга скачана рядом с блокнотом
BOOK = DATA / "zerno_finance.xlsx"
MONTHS = ["Янв", "Фев", "Мар", "Апр", "Май", "Июн", "Июл", "Авг", "Сен"]

try:
    import openpyxl
    print("pandas", pd.__version__, "| openpyxl", openpyxl.__version__, "| книга:", BOOK)
except ImportError:
    print("Нет openpyxl. Выполните в терминале: uv add openpyxl — и перезапустите ядро.")'''),
    ('md', '''## 1. Какие листы в книге

`sheet_name=None` читает все листы сразу и возвращает словарь «имя листа → таблица». Удобно, чтобы осмотреться.'''),
    ('code', '''book = pd.read_excel(BOOK, sheet_name=None)
{name: df.shape for name, df in book.items()}'''),
    ('md', '''## 2. Лист «План»: пропустить шапку

Сначала — как есть. Заголовок листа стал названием первого столбца, остальные — `Unnamed`.'''),
    ('code', '''raw = pd.read_excel(BOOK, sheet_name="План")
raw.head(4)'''),
    ('code', '''plan = pd.read_excel(BOOK, sheet_name="План", skiprows=3)
print(plan.dtypes.value_counts())
plan'''),
    ('md', '''Названия столбцов на месте, месяцы — числа. В отличие от CSV, `decimal=","` не нужен: Excel хранит числа как числа, запятая — только способ их показать.

## 3. Лист «Факт»: шапка выше, столбцы буквами

У «Факта» четыре строки над таблицей. Столбцы выбираем буквами, как при выделении в Excel: A — код, D:L — январь–сентябрь.'''),
    ('code', '''fact = pd.read_excel(BOOK, sheet_name="Факт", skiprows=4, usecols="A,D:L")
fact'''),
    ('md', '''## 4. Строка «Итого»: убрать и сверить

Сумма по кофейням должна совпасть с итоговой строкой. `assert` остановит блокнот, если это не так.'''),
    ('code', '''body = fact[fact["Код"] != "Итого"]
total = fact[fact["Код"] == "Итого"]
by_shops = round(body[MONTHS].sum().sum(), 1)
from_total = round(total[MONTHS].sum().sum(), 1)
assert abs(by_shops - from_total) < 0.5, "сумма по кофейням не сходится с «Итого»"
print("факт за 9 месяцев:", by_shops, "тыс. ₽ — сходится с «Итого»")'''),
    ('md', '''## 5. План против факта'''),
    ('code', '''p = plan[plan["Код"] != "Итого"].set_index("Код")[MONTHS].sum(axis=1)
f = body.set_index("Код")[MONTHS].sum(axis=1)
done = (f / p * 100).round(1)[p > 0].sort_values()
done'''),
    ('md', '''## 6. Сохранить отчёт обратно в Excel

`pd.ExcelWriter` пишет несколько таблиц на разные листы одной книги. Файл появится рядом с блокнотом.'''),
    ('code', '''report = pd.DataFrame({"план": p, "факт": f, "выполнение, %": (f / p * 100).round(1)})[p > 0]
with pd.ExcelWriter("plan_fact_report.xlsx") as writer:
    report.to_excel(writer, sheet_name="Выполнение")
    body.to_excel(writer, sheet_name="Факт без итога", index=False)
list(pd.read_excel("plan_fact_report.xlsx", sheet_name=None))'''),
    ('md', '''## Упражнения

**1.** Прочитайте лист «Факт» не через `skiprows`, а через `header=` — так, чтобы получилась та же таблица `fact` (те же столбцы A, D:L). Результат — в `fact_h`.'''),
    ('code', '''fact_h = None   # ✎ ваш код'''),
    ('code', '''if fact_h is None:
    print("✎ упражнение 1 ещё не решено")
else:
    assert fact_h.equals(fact), "таблицы различаются: проверьте номер строки заголовка (считается с нуля)"
    print("✓ упражнение 1")'''),
    ('md', '''**2.** Прочитайте **все** листы книги одним вызовом и сохраните в `sheets` список их имён в том порядке, в каком они идут в книге.'''),
    ('code', '''sheets = None   # ✎ ваш код'''),
    ('code', '''if sheets is None:
    print("✎ упражнение 2 ещё не решено")
else:
    assert list(sheets) == ["Описание", "План", "Факт"], f"получилось {sheets}"
    print("✓ упражнение 2")'''),
    ('md', '''**3.** В каких месяцах сеть целиком (сумма по кофейням) выполнила план меньше чем на 97 %? Сохраните в `weak` список названий месяцев («Янв», «Фев», …).'''),
    ('code', '''weak = None   # ✎ ваш код'''),
    ('code', '''if weak is None:
    print("✎ упражнение 3 ещё не решено")
else:
    pm = plan[plan["Код"] != "Итого"][MONTHS].sum()
    fm = body[MONTHS].sum()
    want = list((fm / pm * 100)[lambda s: s < 97].index)
    assert list(weak) == want, f"получилось {list(weak)}"
    print("✓ упражнение 3:", want)'''),
]


def write_notebook(path):
    import json

    cells = []
    for i, (kind, src) in enumerate(NB_CELLS):
        cell = {'cell_type': 'markdown' if kind == 'md' else 'code', 'id': f'ch03-excel-{i:02d}',
                'metadata': {}, 'source': src.splitlines(keepends=True)}
        if kind == 'code':
            cell.update(execution_count=None, outputs=[])
        cells.append(cell)
    nb = {'cells': cells, 'nbformat': 4, 'nbformat_minor': 5,
          'metadata': {'kernelspec': {'display_name': 'Python 3', 'language': 'python', 'name': 'python3'},
                       'language_info': {'name': 'python'}}}
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(nb, ensure_ascii=False, indent=1) + '\n', encoding='utf-8')


def main():
    plan, fact = build()
    write_csv(plan, OUT / 'finance_plan.csv')
    write_csv(fact, OUT / 'finance_fact.csv')
    write_xlsx(plan, fact, OUT / 'zerno_finance.xlsx')
    write_notebook(NB_OUT)
    print('ok: zerno_finance.xlsx, finance_plan.csv, finance_fact.csv, notebooks/extra/ch03-excel.ipynb')


if __name__ == '__main__':
    main()
